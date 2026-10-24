import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(14)
    title_run = title_p.add_run("EVENT PROJECT SUBMISSION DOSSIER")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

    # Add subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(20)
    sub_run = sub_p.add_run("Official Proposal & Technical Briefing | ResQAI Platform")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

    # Read the markdown file
    md_path = r"c:\Users\Akshaya\Downloads\ideaverse\ResQAI\PROJECT_SUBMISSION_DOCUMENT.md"
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()

    # Split into lines
    lines = text.split("\n")
    i = 0
    total = len(lines)

    # The exact field headers to detect
    headers = [
        "Idea title:",
        "Theme:",
        "Problem statement:",
        "What problem are you solving (500 words)",
        "What is your proposed solution (1000 words)",
        "Who is affected, and where does the problem occur (300 words)",
        "What is innovative or different about your solution (300 words)",
        "How does your solution work (300 words)",
        "What stage is your solution at",
        "Who will benefit and what impact will it create (300 words)",
        "What is the estimated cost of your solution (300 words)",
        "How can the solution be implemented and scaled (300 words)",
        "What evidence you have",
        "Describe your evidence (300 words)",
        "Who could adopt or implement the solution (300 words)",
        "What support do you need to take the solution forward (300 words)"
    ]

    while i < total:
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        # Check if line matches a header
        matched_header = None
        for h in headers:
            if line.lower() == h.lower():
                matched_header = h
                break

        if matched_header:
            # Add a styled heading box / title
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(16)
            h_p.paragraph_format.space_after = Pt(6)
            h_p.paragraph_format.keep_with_next = True
            
            h_run = h_p.add_run(matched_header)
            h_run.font.name = "Arial"
            h_run.font.size = Pt(13)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(29, 78, 216) # Blue 700

            i += 1
            continue

        # Check if line is a choice row
        if line.startswith("Idea/ Prototype/") or line.startswith("Prototype/ Pilot/"):
            choice_p = doc.add_paragraph()
            choice_p.paragraph_format.space_before = Pt(2)
            choice_p.paragraph_format.space_after = Pt(6)
            c_run = choice_p.add_run(line)
            c_run.font.name = "Calibri"
            c_run.font.size = Pt(10.5)
            c_run.font.italic = True
            c_run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Check if line is bullet / numbered
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content_text = line[2:].strip()
            
            # Format bold prefixes if any (e.g. "Vulnerable Citizens: ...")
            if ":" in content_text and len(content_text.split(":")[0]) < 40:
                prefix, rest = content_text.split(":", 1)
                r1 = p.add_run(prefix + ":")
                r1.font.name = "Calibri"
                r1.font.size = Pt(11)
                r1.font.bold = True
                r2 = p.add_run(rest)
                r2.font.name = "Calibri"
                r2.font.size = Pt(11)
            else:
                r = p.add_run(content_text)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
            i += 1
            continue

        # Numbered list or subsection like "1. Affected Demographics:" or "Phase 1: ..."
        if (len(line) > 2 and line[0].isdigit() and line[1] in [".", ")"]) or line.startswith("Phase "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line)
            r.font.name = "Calibri"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        # Check if it's the selected choice (e.g., "Working prototype" or "Prototype, testing")
        if line in ["Working prototype", "Prototype, testing"]:
            r = p.add_run("Selected Option: ")
            r.font.name = "Calibri"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green
            
            r2 = p.add_run(line)
            r2.font.name = "Calibri"
            r2.font.size = Pt(11.5)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor(15, 23, 42)
        else:
            r = p.add_run(line)
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

        i += 1

    out_path = r"c:\Users\Akshaya\Downloads\ideaverse\ResQAI\ResQAI_Submission_Document.docx"
    doc.save(out_path)
    print(f"Document successfully created at {out_path} ({os.path.getsize(out_path)} bytes)")

if __name__ == "__main__":
    create_document()
